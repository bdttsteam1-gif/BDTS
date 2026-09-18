#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
convert.py — BDTS-main(구글 시트 허브) 내려받기 xlsx → bdts-portal-main(Azure) 저장 형식 JSON

사용법:
  python3 convert.py 거래처이력카드.xlsx [--out seed] [--base ../../data/domestic-base.js]

입력  : 구글 시트 백업본(거래처이력카드.xlsx). 탭 = records / claim_progress / claim_extra /
        csc_reference_log / cartridge_log (그 외 탭은 무시)
출력  : <out>/client_cards.json, csc_cards.json, claim_extra.json, csc_reference_log.json,
        cartridge_log.json  +  <out>/이관_검토리포트.md  +  <out>/manifest.json

원칙
  * 원본 값은 바꾸지 않습니다. 날짜는 화면이 쓰는 'YYYY-MM-DD' 로 표기만 통일하고,
    표기를 바꾼 건 전부 리포트에 남깁니다(무엇을 무엇으로 왜).
  * 해석할 수 없는 날짜는 비워 두고 원본 문자열을 _raw* 에 보존합니다(임의 추정 금지).
  * 목록(장비·아이템·이름)에 없는 값은 고치지 않고 그대로 두되 리포트에 모아 보여줍니다.
  * 마스터로그(클레임 로그 국내)는 만들지 않습니다(logs: []). 과거 클레임은 이미
    data/domestic-rows.js 씨앗에 들어 있어 이관하면 중복이 됩니다.
"""
import sys, os, re, json, argparse, datetime, collections

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl 이 필요합니다:  pip install openpyxl')

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------- 옵션
ap = argparse.ArgumentParser()
ap.add_argument('xlsx')
ap.add_argument('--out', default=os.path.join(HERE, 'seed'))
ap.add_argument("--base", default=os.path.join(HERE, "..", "..", "data", "domestic-base.js"),
                help='포털 기준 목록(장비·아이템·이름·병원)이 든 domestic-base.js')
ap.add_argument('--title-from', choices=['request', 'none'], default='request',
                help='거래처 이력카드 Title(포털 신규 항목)을 무엇으로 채울지: request=접수사항 첫 줄')
ap.add_argument('--derive-hospital', action='store_true', default=True,
                help='CSC 병원명이 비어 있고 제목이 "병원 / …" 꼴이면 앞부분을 병원명으로(목록에 있을 때만)')
ap.add_argument('--no-derive-hospital', dest='derive_hospital', action='store_false')
ap.add_argument('--value-map', default=os.path.join(HERE, 'value-map.json'),
                help='사람이 확정한 값 정리표 {"화면": {"항목": {"원본": "바꿀 값"}}} — 적용 내역은 리포트와 각 건의 _valueFix 에')
ap.add_argument('--overrides', default=os.path.join(HERE, 'overrides.json'),
                help='사람이 확인한 값 보정표 {"시트": {"id": {"항목": "값"}}} — 적용 내역은 리포트에 남습니다')
args = ap.parse_args()

os.makedirs(args.out, exist_ok=True)
OVR = {}
if os.path.exists(args.overrides):
    try: OVR = json.load(open(args.overrides, encoding='utf-8'))
    except Exception as e: sys.exit('overrides.json 을 읽지 못했습니다: %s' % e)
VMAP = {}
if os.path.exists(args.value_map):
    try: VMAP = json.load(open(args.value_map, encoding='utf-8'))
    except Exception as e: sys.exit('value-map.json 을 읽지 못했습니다: %s' % e)
VMAP_USED = collections.Counter()   # (화면, 항목, 원본, 바꿀 값) → 건수
def vfix(rec, sheet, cat, field, val):
    """값 정리표에 있으면 바꾸고, 바꾼 내역을 건(_valueFix)과 집계에 남깁니다. 없으면 원본 그대로."""
    new = ((VMAP.get(sheet) or {}).get(cat) or {}).get(val)
    if new is None or new == val: return val
    rec.setdefault('_valueFix', []).append({'field': field, 'from': val, 'to': new})
    VMAP_USED[(sheet, cat, val, new)] += 1
    return new
OVR_USED = []   # (시트, id, 항목, 원본, 값)
def override(sheet, rid, field, cur, raw=None):
    v = (OVR.get(sheet) or {}).get(str(rid), {})
    if field in v:
        OVR_USED.append((sheet, str(rid), field, raw if raw is not None else cur, v[field]))
        return v[field]
    return cur
REPORT = []            # 리포트 줄들
def rep(s=''): REPORT.append(s)

# ---------------------------------------------------------------- 기준 목록
BASE = {'device': [], 'item': [], 'names': [], 'hospitals': []}
if os.path.exists(args.base):
    b = open(args.base, encoding='utf-8').read()
    BASE = json.loads(b[b.find('{'):b.rfind('}') + 1])
DEVICES = set(BASE.get('device', []))
ITEMS = set(BASE.get('item', []))
NAMES_KO = set(n['ko'] for n in BASE.get('names', []))
HOSPS = set(h['name'] for h in BASE.get('hospitals', []))

# ---------------------------------------------------------------- 공통 유틸
def S(v):
    """셀 값 → 문자열(None→'', 숫자 정수화, 엑셀 오류값(#VALUE! 등)→'')"""
    if v is None: return ''
    if isinstance(v, str) and re.match(r'^#(VALUE|REF|N/A|DIV/0|NAME|NUM|NULL)[!?]?$', v.strip()): return ''
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
    return str(v).strip()

DATE_FIX = []   # (시트, id, 항목, 원본, 결과, 사유)
DATE_BAD = []   # (시트, id, 항목, 원본)

def norm_date(v, sheet, rid, field):
    """날짜 → 'YYYY-MM-DD'. 못 읽으면 '' + DATE_BAD 기록."""
    if v is None or v == '': return ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})(?:[ T]\d{2}:\d{2}(?::\d{2})?(?:\.\d+)?Z?)?$', s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= mo <= 12 and 1 <= d <= 31:
            out = '%04d-%02d-%02d' % (y, mo, d)
            if out != s: DATE_FIX.append((sheet, rid, field, s, out, '시각 잘라냄'))
            return out
        DATE_BAD.append((sheet, rid, field, s)); return ''
    # 2024.10.11 / 2025--08-21 / 2024.05-13 / 2024/10/11
    m = re.match(r'^(\d{4})[.\-/]+(\d{1,2})[.\-/]+(\d{1,2})\.?$', s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= mo <= 12 and 1 <= d <= 31:
            out = '%04d-%02d-%02d' % (y, mo, d)
            DATE_FIX.append((sheet, rid, field, s, out, '구분자 통일'))
            return out
    DATE_BAD.append((sheet, rid, field, s)); return ''

def norm_datetime(v, sheet, rid, field):
    """'2026-05-06 10:11' / datetime → ('YYYY-MM-DD', 'HH:MM')"""
    if v is None or v == '': return '', ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d'), v.strftime('%H:%M')
    s = str(v).strip()
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})[ T](\d{2}):(\d{2})', s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= mo <= 12 and 1 <= d <= 31:
            return '%04d-%02d-%02d' % (y, mo, d), '%s:%s' % (m.group(4), m.group(5))
        DATE_BAD.append((sheet, rid, field, s)); return '', ''
    d = norm_date(s, sheet, rid, field)
    return d, ''

def fmt_phone(s):
    """포털 fmtPhone 과 같은 규칙: 숫자만 남겨 000-0000-0000 (02 는 2자리, 15xx 는 4자리). 뒷글자(내선 등) 유지"""
    s = S(s)
    if not s: return ''
    m = re.match(r'^([\d\-\s\.\(\)]+)(.*)$', s)
    if not m: return s
    digits = re.sub(r'\D', '', m.group(1)); tail = m.group(2).strip()
    if not digits: return s
    if digits.startswith('02'):
        p = ['02', digits[2:-4], digits[-4:]] if len(digits) >= 9 else ['02', digits[2:]]
    elif re.match(r'^1[5-9]\d\d', digits) and len(digits) == 8:
        p = [digits[:4], digits[4:]]
    elif len(digits) >= 10:
        p = [digits[:3], digits[3:-4], digits[-4:]]
    elif len(digits) >= 8:
        p = [digits[:-4], digits[-4:]]
    else:
        p = [digits]
    out = '-'.join(x for x in p if x)
    return (out + (' ' + tail if tail else '')).strip()

def first_line(s, n=80):
    s = S(s).replace('\r', '')
    line = s.split('\n', 1)[0].strip()
    return (line[:n] + '…') if len(line) > n else line

def load_sheet(wb, name):
    if name not in wb.sheetnames: return [], []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return [], []
    hdr = [S(h) for h in rows[0]]
    data = [r for r in rows[1:] if any(c not in (None, '') for c in r)]
    return hdr, data

def rowdict(hdr, r):
    d = {}
    for i, h in enumerate(hdr):
        if h: d[h] = r[i] if i < len(r) else None
    return d

def dedupe(rows, key='id', label=''):
    """같은 id 가 여러 번이면 첫 건만. 내용이 다르면 리포트에 남김."""
    seen, out, same, diff = {}, [], 0, []
    for r in rows:
        k = S(r.get(key))
        if k in seen:
            if json.dumps(seen[k], sort_keys=True, ensure_ascii=False, default=str) == json.dumps(r, sort_keys=True, ensure_ascii=False, default=str):
                same += 1
            else:
                diff.append(k)
            continue
        seen[k] = r; out.append(r)
    if same or diff:
        rep('- %s: 같은 id 중복 — 완전 동일 %d건 제거%s' % (label, same,
            (', **내용이 다른 중복 %d건**(첫 건만 남김, 확인 필요): %s' % (len(diff), ', '.join(diff))) if diff else ''))
    return out

# ---------------------------------------------------------------- 읽기
wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
rep('# BDTS-main → bdts-portal-main 이관 검토 리포트')
rep('')
rep('- 원본: `%s`  (탭: %s)' % (os.path.basename(args.xlsx), ', '.join(wb.sheetnames)))
rep('- 생성: %s' % datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))
rep('- 기준 목록: `%s` (장비 %d · 아이템 %d · 이름 %d · 병원 %d)' % (os.path.basename(args.base), len(DEVICES), len(ITEMS), len(NAMES_KO), len(HOSPS)))
rep('')

manifest = {}

# ================================================================ 1. 거래처 이력카드
hdr, data = load_sheet(wb, 'records')
raw = [rowdict(hdr, r) for r in data]
rep('## 1. 거래처 이력카드 (records → client_cards)')
rep('')
rep('- 원본 행: %d건 (빈 줄 제외)' % len(raw))
raw = dedupe(raw, label='records')

# 원본 id → 정수/문자
def norm_id(v):
    s = S(v)
    return int(s) if re.match(r'^\d+$', s) else s

cards, notlist = [], collections.defaultdict(collections.Counter)
NL_IDS = collections.defaultdict(lambda: collections.defaultdict(list))
def nl(sheet, cat, val, rid):
    notlist_ref[cat][val] += 1; NL_IDS[(sheet, cat)][val].append(str(rid))
CLIENT_HOSPS = set()
for r in raw:
    rid = norm_id(r.get('id'))
    date = norm_date(r.get('date'), 'records', rid, 'date')
    date = override('records', rid, 'date', date, S(r.get('date')))
    if not date: DATE_BAD.append(('records', rid, 'date', S(r.get('date')) or '(빈 값)'))
    reagent, lot = S(r.get('reagent')), S(r.get('reagentLot'))
    items = [{'item': reagent, 'lot': lot}] if (reagent or lot) else []
    st = {'device': S(r.get('device')), 'sn': S(r.get('sn')), 'sw': S(r.get('sw')),
          'items': items, 'req': S(r.get('request')), 'act': S(r.get('action')),
          'etc': S(r.get('etc')),      # BDTS-main "기타 (특이사항, 현황, 정보)" — 기기 블록 안 항목. 포털 입력 화면에는 아직 없음
          'c1': '', 'c2': '', 'c3': '', 'files': []}
    # 예전 이미지(구글 드라이브 URL 목록)는 첨부로 보이게, 원본도 보존
    images = S(r.get('images'))
    try:
        il = json.loads(images) if images else []
    except Exception:
        il = []
    for k, u in enumerate([x for x in il if isinstance(x, str) and x]):
        st['files'].append({'name': '사진 %d (Google Drive)' % (k + 1), 'data': u, 'legacyUrl': True})
    card = {
        'id': rid,
        'manager': S(r.get('manager')), 'contact': S(r.get('contact')), 'date': date,
        'hospital': S(r.get('hospital')), 'requester': S(r.get('requester')),
        'region': S(r.get('region')), 'phone': fmt_phone(r.get('phone')), 'agency': S(r.get('agency')),
        'category': S(r.get('category')),
        'title': first_line(r.get('request')) if args.title_from == 'request' else '',
        'doctorNote': S(r.get('doctorNote')),
        'folder': S(r.get('folder')),                 # 원본 파일 폴더 (통합 검색 상세에 표시)
        'sets': [st], 'logs': [], 'logMap': [],
        '_src': 'BDTS-main', '_srcId': S(r.get('id')),
    }
    if not date and r.get('date') not in (None, ''):
        card['_rawDate'] = S(r.get('date'))
    # 파생(포털 clSyncDerived 와 동일): devices(로그 1건=1행) / cls / files
    its = items or [{'item': '', 'lot': ''}]
    card['devices'] = [{'device': st['device'], 'sn': st['sn'], 'sw': st['sw'], 'item': it['item'], 'lot': it['lot'],
                        'req': st['req'], 'act': st['act'], 'etc': st['etc'], 'files': st['files']} for it in its]
    card['cls'] = {'c1': '', 'c2': '', 'c3': ''}
    card['files'] = list(st['files'])
    # 사용자 확정 값 정리표 적용 (원본은 _valueFix 에 남김)
    card['category'] = vfix(card, '거래처 이력카드', '분류', 'category', card['category'])
    card['manager']  = vfix(card, '거래처 이력카드', '담당자', 'manager', card['manager'])
    st['device']     = vfix(card, '거래처 이력카드', '기기명', 'sets[0].device', st['device'])
    if st['items']:
        st['items'][0]['item'] = vfix(card, '거래처 이력카드', '시약명', 'sets[0].items[0].item', st['items'][0]['item'])
    reagent = st['items'][0]['item'] if st['items'] else ''
    for dv in card['devices']: dv['device'] = st['device']; dv['item'] = reagent
    cards.append(card)
    CLIENT_HOSPS.add(card['hospital'])
    notlist_ref = notlist
    if st['device'] and st['device'] not in DEVICES: nl('거래처 이력카드', '기기명', st['device'], rid)
    if reagent and reagent not in ITEMS: nl('거래처 이력카드', '시약명', reagent, rid)
    if card['manager'] and card['manager'] not in NAMES_KO: nl('거래처 이력카드', '담당자', card['manager'], rid)
    if card['category'] and card['category'] not in ['문의', '컴플레인', 'PM(Preventive Maintenance)', '내부 성능 평가', 'AS 수리', '세팅', '기타']:
        nl('거래처 이력카드', '분류', card['category'], rid)
    if card['contact'] and card['contact'] not in ['방문', '전화', '메일', '기타']:
        nl('거래처 이력카드', '컨택방식', card['contact'], rid)

cards.sort(key=lambda c: (c['date'] or '0000-00-00', str(c['id'])), reverse=True)
json.dump(cards, open(os.path.join(args.out, 'client_cards.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
manifest['client_cards'] = len(cards)
rep('- 변환 결과: **%d건** → `client_cards.json`' % len(cards))
rep('- 형식: 포털 신규 입력과 같은 구조(`sets[1]` = 기기+시약+접수/조치, `logs: []`). Title 은 %s' %
    ('접수사항 첫 줄(80자)' if args.title_from == 'request' else '비워 둠'))
rep('- 포털 입력 화면에 없는 항목 보존: 기기 블록의 `etc`(기타 특이사항·현황·정보 %d건 → 상세 팝업·검색에 표시), `folder`(원본 폴더 %d건), 구글드라이브 사진 URL %d건' %
    (sum(1 for c in cards if c['sets'][0]['etc']), sum(1 for c in cards if c['folder']), sum(len(c['files']) for c in cards)))
rep('- 병원 수: %d · 담당자: %s' % (len(CLIENT_HOSPS), ', '.join('%s %d' % kv for kv in collections.Counter(c['manager'] for c in cards).most_common(6))))
rep('')
for k, cnt in notlist.items():
    rep('- 목록에 없는 %s 값(그대로 둠): %s' % (k, ', '.join('`%s` %d' % kv for kv in cnt.most_common(30))))
rep('')

# ================================================================ 2. 고객지원센터
hdr, data = load_sheet(wb, 'claim_progress')
raw = [rowdict(hdr, r) for r in data]
rep('## 2. 고객지원센터 (claim_progress → csc_cards)')
rep('')
rep('- 원본 행: %d건' % len(raw))
raw = dedupe(raw, label='claim_progress')

VIA_MAP = {'메일': 'email', 'CSC': 'CSC(고객지원센터)', 'QR': 'QR Code', '전화': 'phone call', 'SNS': 'SNS', '방문': 'Visit', 'AFIAS NET': 'AFIAS NET'}
ACT_MAP = {'전화': 'Phone call response', '메일': 'Email response', 'CSC': 'CSC(고객지원센터)', 'SNS': 'SNS response', '방문': 'Visit', 'AFIAS NET': 'AFIAS NET', 'QR': 'QR Code'}
KNOWN_HOSPS = HOSPS | CLIENT_HOSPS | set(S(r.get('hospital')) for r in raw if S(r.get('hospital')))

def jload(v):
    s = S(v)
    if not s or s == '[]': return []
    try:
        x = json.loads(s)
        return x if isinstance(x, list) else []
    except Exception:
        return []

cscs, derived_h, unknown_via, unknown_act, nl2 = [], [], collections.Counter(), collections.Counter(), collections.defaultdict(collections.Counter)
for r in raw:
    rid = S(r.get('id'))
    d, t = norm_datetime(r.get('recvAt'), 'claim_progress', rid, 'recvAt')
    rd, rt = norm_datetime(r.get('replyAt'), 'claim_progress', rid, 'replyAt')
    d = override('claim_progress', rid, 'date', d, S(r.get('recvAt')))
    rd = override('claim_progress', rid, 'replied', rd, S(r.get('replyAt')))
    if rd and not rt:   # 보정표로 날짜만 확정한 경우, 원본에 적힌 시각(HH:MM)은 그대로 살립니다
        _m = re.search(r'(\d{2}):(\d{2})', S(r.get('replyAt'))); rt = _m.group(0) if _m else ''
    if not d: DATE_BAD.append(('claim_progress', rid, 'recvAt', S(r.get('recvAt')) or '(빈 값)'))
    title, hosp = S(r.get('title')), S(r.get('hospital'))
    if not hosp and args.derive_hospital and ' / ' in title:
        head = title.split(' / ', 1)[0].strip()
        if head in KNOWN_HOSPS:
            hosp = head; derived_h.append((rid, head))
    via_raw, act_raw = S(r.get('via')), S(r.get('replyVia'))
    via = VIA_MAP.get(via_raw, via_raw); act = ACT_MAP.get(act_raw, act_raw)
    if via_raw and via_raw not in VIA_MAP: unknown_via[via_raw] += 1
    if act_raw and act_raw not in ACT_MAP: unknown_act[act_raw] += 1
    status_raw = S(r.get('status'))
    done = status_raw in ('완료', '완료됨', 'done')
    # 장비 세트 — devicesJson / itemsJson 이 있으면 그것, 없으면 낱개 칸
    devs = jload(r.get('devicesJson'))
    itms = jload(r.get('itemsJson'))
    if not devs:
        devs = [{'device': S(r.get('device')), 'sample': S(r.get('sample')), 'sampleOther': S(r.get('sampleOther')), 'store': S(r.get('store'))}] \
               if any(S(r.get(k)) for k in ('device', 'serial', 'sw', 'sample', 'store')) else []
    if not itms:
        if any(S(r.get(k)) for k in ('item', 'itemLot', 'ctrlLot', 'itemUnit', 'anticoag')):
            itms = [{'item': S(r.get('item')), 'itemLot': S(r.get('itemLot')), 'unit': S(r.get('itemUnit')),
                     'ctrlItem': S(r.get('ctrlItem')), 'ctrlLot': S(r.get('ctrlLot')), 'anticoag': S(r.get('anticoag'))}]
    d0 = devs[0] if devs else {}
    def mk_item(x):
        smp = S(x.get('sample')) or S(d0.get('sample')) or S(r.get('sample'))
        if smp in ('Other', '기타', 'other'):
            smp = S(x.get('sampleOther')) or S(d0.get('sampleOther')) or S(r.get('sampleOther')) or smp
        return {'item': S(x.get('item')), 'lot': S(x.get('itemLot') or x.get('lot')), 'ctrl': S(x.get('ctrlLot') or x.get('ctrl')),
                'unit': S(x.get('unit')), 'sample': smp,
                'store': S(x.get('store')) or S(d0.get('store')) or S(r.get('store')),
                'anti': S(x.get('anticoag') or x.get('anti'))}
    items = [it for it in (mk_item(x) for x in itms) if any(it.values())]
    # 사용자 확정 값 정리표 + "TSH, T3" 처럼 쉼표로 묶인 아이템은 아이템 여러 개로 나눕니다(LOT 등은 첫 아이템에)
    _rec = {}
    split_items = []
    for it in items:
        it['item'] = vfix(_rec, '고객지원센터', '아이템', 'sets[0].items[].item', it['item'])
        parts = [p.strip() for p in it['item'].split(',')] if ',' in it['item'] else [it['item']]
        if len(parts) > 1:
            _rec.setdefault('_valueFix', []).append({'field': 'sets[0].items[].item', 'from': it['item'], 'to': ' / '.join(parts), 'note': '아이템 %d개로 나눔' % len(parts)})
            VMAP_USED[('고객지원센터', '아이템', it['item'], '아이템 %d개로 나눔' % len(parts))] += 1
        for k, pn in enumerate(parts):
            split_items.append(dict(it, item=pn) if k == 0 else {'item': pn, 'lot': '', 'ctrl': '', 'unit': '', 'sample': it['sample'], 'store': it['store'], 'anti': it['anti']})
    items = split_items
    for dv in devs: dv['device'] = vfix(_rec, '고객지원센터', '장비', 'sets[].device', S(dv.get('device')))
    sets = []
    for i, dv in enumerate(devs or [{}]):
        sets.append({'device': S(dv.get('device')), 'sn': S(r.get('serial')) if i == 0 else '', 'sw': S(r.get('sw')) if i == 0 else '',
                     'items': items if i == 0 else [], 'issue': S(r.get('complaint')) if i == 0 else '',
                     'c1': '', 'c2': '', 'c3': '', 'files': []})
    phone = fmt_phone(r.get('phone'))
    card = {
        'id': rid,
        'date': d, 'time': t, 'receiver': vfix(_rec, '고객지원센터', '접수담당자', 'receiver', S(r.get('writer'))),
        'via': via, 'title': title, 'kind': S(r.get('kind')) or 'TS',
        'hospital': hosp, 'region': S(r.get('region')),
        'phones': [{'no': phone, 'memo': ''}] if phone else [], 'phone': phone,
        'person': S(r.get('contact')),
        'sets': sets,
        'confirm': vfix(_rec, '고객지원센터', '회신자', 'confirm', S(r.get('replier'))), 'replied': rd, 'rtime': rt,
        'action': act, 'status': '완료' if done else '진행 중', 'end': rd if done else '',
        'note': S(r.get('note')),
        'mailSubject': S(r.get('mailSubject')), 'cscKeyword': S(r.get('cscKeyword')),
        'logs': [], 'logMap': [],
        '_src': 'BDTS-main', '_srcId': rid,
    }
    if _rec.get('_valueFix'): card['_valueFix'] = _rec['_valueFix']
    if not d and r.get('recvAt') not in (None, ''): card['_rawRecvAt'] = S(r.get('recvAt'))
    if not rd and r.get('replyAt') not in (None, ''): card['_rawReplyAt'] = S(r.get('replyAt'))
    if done and not rd:
        # 사용자 결정(2026-09-09): 회신일이 없으면 완료로 볼 수 없다 → On going 으로 두고,
        # 클레임 파트에서 END Date 를 적어 Closed 로 바꾼다. 날짜를 임의로 채우지 않는다.
        card['status'] = '진행 중'; card['end'] = ''; card['_noReplyDate'] = True
    # 파생(포털 csSyncDerived 와 동일)
    card['devices'] = [{'device': s_['device'], 'sn': s_['sn'], 'sw': s_['sw']} for s_ in sets]
    card['items'] = list(items)
    card['issue'] = sets[0]['issue']
    card['files'] = []
    card['c1'] = card['c2'] = card['c3'] = ''
    cscs.append(card)
    notlist_ref = nl2
    for s_ in sets:
        if s_['device'] and s_['device'] not in DEVICES: nl('고객지원센터', '장비', s_['device'], rid)
    for it in items:
        if it['item'] and it['item'] not in ITEMS: nl('고객지원센터', '아이템', it['item'], rid)
    for nm, lab in ((card['receiver'], '접수담당자'), (card['confirm'], '회신자')):
        if nm and nm not in NAMES_KO: nl('고객지원센터', lab, nm, rid)
    if act_raw and act_raw not in ACT_MAP: nl('고객지원센터', '회신 경로', act_raw, rid)

cscs.sort(key=lambda c: (c['date'] or '0000-00-00', c['time'], c['id']), reverse=True)
json.dump(cscs, open(os.path.join(args.out, 'csc_cards.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
manifest['csc_cards'] = len(cscs)
rep('- 변환 결과: **%d건** → `csc_cards.json` (완료 %d · 진행 중 %d · 종류 그외 %d)' %
    (len(cscs), sum(1 for c in cscs if c['status'] == '완료'), sum(1 for c in cscs if c['status'] != '완료'), sum(1 for c in cscs if c['kind'] == '그외')))
rep('- 접수 경로 → Via 변환: ' + ', '.join('%s→%s' % kv for kv in VIA_MAP.items() if kv[0] in set(S(x.get('via')) for x in raw)))
rep('- 회신 경로 → Action 변환: ' + ', '.join('%s→%s' % kv for kv in ACT_MAP.items() if kv[0] in set(S(x.get('replyVia')) for x in raw)))
if unknown_via: rep('- 변환표에 없는 접수 경로(그대로 둠): ' + ', '.join('`%s` %d' % kv for kv in unknown_via.items()))
if unknown_act: rep('- 변환표에 없는 회신 경로(그대로 둠): ' + ', '.join('`%s` %d' % kv for kv in unknown_act.items()))
rep('- 병원명이 비어 있던 건: %d건 → 제목 "병원 / …" 앞부분이 병원 목록에 있어 채운 건 **%d건**, 여전히 빈 건 %d건' %
    (sum(1 for x in raw if not S(x.get('hospital'))), len(derived_h), sum(1 for c in cscs if not c['hospital'])))
nr = [c for c in cscs if c.get('_noReplyDate')]
rep('- 원본은 "완료"인데 회신일이 없는 건: **%d건 → 진행 중(On going)으로 이관** (사용자 결정: 회신일 없이는 완료가 아님. 클레임 파트에서 END Date 를 적어 Closed 로 바꿈): %s' %
    (len(nr), ', '.join('%s(%s)' % (c['id'], c['hospital'] or c['title'][:20]) for c in nr)))
for k, cnt in nl2.items():
    rep('- 목록에 없는 %s 값(그대로 둠): %s' % (k, ', '.join('`%s` %d' % kv for kv in cnt.most_common(30))))
rep('')

# ================================================================ 3. 통합 클레임 로그 업로드분
hdr, data = load_sheet(wb, 'claim_extra')
raw = [rowdict(hdr, r) for r in data]
rep('## 3. 통합 클레임 로그 업로드분 (claim_extra → claim_extra)')
rep('')
rep('- 원본 행: %d건' % len(raw))
raw = dedupe(raw, label='claim_extra')
extras = []
for r in raw:
    rid = S(r.get('id'))
    o = {}
    for h in hdr:
        if not h: continue
        v = r.get(h)
        if h in ('recv_date', 'replied'): o[h] = norm_date(v, 'claim_extra', rid, h)
        elif h == 'year': o[h] = S(v)
        elif h == 'dur': o[h] = S(v)
        else: o[h] = S(v)
    extras.append(o)
json.dump(extras, open(os.path.join(args.out, 'claim_extra.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
manifest['claim_extra'] = len(extras)
yrs = collections.Counter(e['year'] for e in extras)
nos = [e['no'] for e in extras if e.get('no')]
rep('- 변환 결과: **%d건** → `claim_extra.json` (연도 %s · 관리번호 %s ~ %s)' %
    (len(extras), ', '.join('%s %d' % kv for kv in sorted(yrs.items())), min(nos) if nos else '-', max(nos) if nos else '-'))
rep('- 포털 씨앗(data/claims-YYYY.json)과의 중복은 화면(claim-sheet.js)이 "관리번호+접수일+제목"으로 걸러 보여줍니다.')
rep('')

# ================================================================ 4. 참고 기록 · 카트리지
for sheet in ('csc_reference_log', 'cartridge_log'):
    hdr, data = load_sheet(wb, sheet)
    raw = dedupe([rowdict(hdr, r) for r in data], label=sheet)
    out = []
    for r in raw:
        rid = S(r.get('id'))
        o = {h: (norm_date(r.get(h), sheet, rid, h) if h == 'date' else S(r.get(h))) for h in hdr if h}
        out.append(o)
    json.dump(out, open(os.path.join(args.out, sheet + '.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    manifest[sheet] = len(out)
rep('## 4. 참고 기록 · 카트리지 기록')
rep('')
rep('- `csc_reference_log.json` %d건 · `cartridge_log.json` %d건 (날짜 표기만 통일)' % (manifest['csc_reference_log'], manifest['cartridge_log']))
rep('- 시트 `upload_blobs` / `csc_faq_extra` / `csc_error_extra` 는 %s' %
    ', '.join('%s %d행' % (n, len(load_sheet(wb, n)[1])) for n in ('upload_blobs', 'csc_faq_extra', 'csc_error_extra') if n in wb.sheetnames) + ' — 포털에서 쓰지 않아 옮기지 않습니다.')
rep('')

# ================================================================ 4-1. 값 정리표 적용 내역
rep('## 4-1. 사용자 확정 값 정리표 적용 (value-map.json)')
rep('')
if VMAP_USED:
    rep('| 화면 | 항목 | 원본 | 바꾼 값 | 건수 |'); rep('|---|---|---|---|---|')
    for (sh, cat, a, b), n in sorted(VMAP_USED.items(), key=lambda kv: (kv[0][0], kv[0][1], -kv[1])):
        rep('| %s | %s | `%s` | `%s` | %d |' % (sh, cat, a, b, n))
    rep(''); rep('- 바꾼 건은 각 기록의 `_valueFix` 에 원본이 남아 있어 추적할 수 있습니다.')
else:
    rep('- 없음')
rep('')

# ================================================================ 5. 날짜 보정 내역
rep('## 5. 날짜 표기 보정 (값이 바뀐 건 전부)')
rep('')
cut = [x for x in DATE_FIX if x[5] == '시각 잘라냄']
fix = [x for x in DATE_FIX if x[5] != '시각 잘라냄']
rep('- `YYYY-MM-DD 00:00:00` 에서 시각만 잘라낸 건: %d건 (날짜 값은 그대로)' % len(cut))
rep('- 구분자·형식을 바꾼 건: **%d건**' % len(fix))
if fix:
    rep(''); rep('| 시트 | id | 항목 | 원본 | 결과 | 사유 |'); rep('|---|---|---|---|---|---|')
    for x in fix: rep('| %s | %s | %s | `%s` | `%s` | %s |' % x)
rep('')
rep('### 사람이 확인한 보정표(overrides.json) 적용 내역')
rep('')
if OVR_USED:
    rep('| 시트 | id | 항목 | 원본 | 적용값 |'); rep('|---|---|---|---|---|')
    for x in OVR_USED: rep('| %s | %s | %s | `%s` | `%s` |' % x)
else:
    rep('- 없음')
rep('')
rep('### 읽지 못해 비워 둔 날짜 (원본은 `_raw…` 항목에 보존 — 확인 후 overrides.json 에 적거나 화면에서 직접 수정)')
rep('')
_seen = set(); DATE_BAD = [x for x in DATE_BAD if not (x in _seen or _seen.add(x))]
_ov = set((a, b, c) for a, b, c, _, _ in OVR_USED)
DATE_BAD = [x for x in DATE_BAD if (x[0], str(x[1]), {'recvAt': 'date', 'replyAt': 'replied'}.get(x[2], x[2])) not in _ov]
if DATE_BAD:
    rep('| 시트 | id | 항목 | 원본 |'); rep('|---|---|---|---|')
    for x in DATE_BAD: rep('| %s | %s | %s | `%s` |' % x)
else:
    rep('- 없음')
rep('')

# ================================================================ 목록에 없는 값 xlsx
import difflib
xw = openpyxl.Workbook(); ws0 = xw.active; ws0.title = '안내'
ws0.append(['이 파일은 이관 시 포털 목록(장비·아이템·이름 등)에 없어 그대로 옮긴 값들입니다.'])
ws0.append(['값은 고치지 않았습니다. 정리 규칙을 정해 주시면 overrides 또는 변환 규칙으로 반영합니다.'])
ws0.append(['"비슷한 목록 값"은 참고용 자동 추천이며 확정이 아닙니다.'])
for (sheet, cat), vals in NL_IDS.items():
    pool = {'기기명': DEVICES, '장비': DEVICES, '시약명': ITEMS, '아이템': ITEMS, '담당자': NAMES_KO, '접수담당자': NAMES_KO, '회신자': NAMES_KO,
            '분류': set(['문의', '컴플레인', 'PM(Preventive Maintenance)', '내부 성능 평가', 'AS 수리', '세팅', '기타']),
            '컨택방식': set(['방문', '전화', '메일', '기타']), '회신 경로': set(ACT_MAP.values())}.get(cat, set())
    w = xw.create_sheet((sheet[:6] + '_' + cat)[:31])
    w.append(['화면', '항목', '값(원본 그대로)', '건수', '비슷한 목록 값(자동 추천)', '해당 id'])
    for v, ids in sorted(vals.items(), key=lambda kv: -len(kv[1])):
        sim = difflib.get_close_matches(v, list(pool), n=3, cutoff=0.5) if pool else []
        w.append([sheet, cat, v, len(ids), ', '.join(sim), ', '.join(ids[:200])])
    for col, wd in zip('ABCDEF', (14, 12, 34, 8, 40, 60)): w.column_dimensions[col].width = wd
xw.save(os.path.join(args.out, '목록에_없는_값.xlsx'))
rep('## 6. 목록에 없는 값')
rep('')
rep('- 값은 고치지 않고 그대로 옮겼습니다. 전체 목록(건수·해당 id·비슷한 목록 값)은 `목록에_없는_값.xlsx` 에 있습니다.')
rep('')

# ================================================================ manifest
manifest['_built'] = datetime.datetime.now().isoformat(timespec='seconds')
manifest['_source'] = os.path.basename(args.xlsx)
json.dump(manifest, open(os.path.join(args.out, 'manifest.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
open(os.path.join(args.out, '이관_검토리포트.md'), 'w', encoding='utf-8').write('\n'.join(REPORT) + '\n')
print('\n'.join(REPORT))
print('\n→ 출력 폴더:', args.out, json.dumps({k: v for k, v in manifest.items() if not k.startswith('_')}, ensure_ascii=False))
